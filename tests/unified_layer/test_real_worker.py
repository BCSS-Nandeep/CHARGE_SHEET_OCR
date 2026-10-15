"""Tests with the REAL, unmodified CHARGE_SHEET_OCR snapshot (its own process) behind the real gateway.

OCR and the vision model are switched off so the runs are fast and deterministic (digital-text PDFs, as in the
snapshot's own tests); real-document accuracy is a separate, opt-in check (see docs/UNIFIED_ARCHITECTURE.md).

The key test compares the unified path with a *plain* CHARGE_SHEET_OCR application (``app.main:app``, no wrapper) given the
same file: the extraction result must be identical, which is what "the adapter does not change extraction" means.
"""

from __future__ import annotations

import copy
import os
import socket
import subprocess
import sys
import time
from pathlib import Path
from types import SimpleNamespace

import httpx
import pytest
from fastapi.testclient import TestClient

from tests.conftest import ADMIN_KEY, OPERATOR_KEY, REVIEWER_KEY
from tests.unified_layer.conftest import wait_for
from tests.unified_layer.helpers import envelope_errors, ir_docx, make_png, report, upload, wait_report
from unified.adapters.worker import WorkerManager
from unified.config import ROOT, load_settings
from unified.factory import create_gateway

pytestmark = pytest.mark.real_worker

VENDOR = ROOT / "vendor" / "charge_sheet_ocr"
SAMPLES = VENDOR / "tests" / "fixtures"
CS_VENV = Path(r"C:\Users\Dell\Desktop\CHARGE_SHEET_OCR\.venv\Scripts\python.exe")
KEYS = f"ops:{OPERATOR_KEY}:operator;rita:{REVIEWER_KEY}:reviewer"
DETERMINISTIC = {"OCR_ENABLED": "false", "VLM_ENABLED": "false", "DOCLING_ENABLED": "false", "VL_ENABLED": "false",
                 "LLM_ENABLED": "false", "ASYNC_PROCESSING": "false", "OFFLINE_MODE": "true",
                 "ALLOW_EXTERNAL_NETWORK": "false"}


def worker_python() -> str:
    explicit = os.environ.get("UNIFIED_TEST_WORKER_PYTHON")
    if explicit:
        return explicit
    return str(CS_VENV) if CS_VENV.is_file() else sys.executable


def free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


class HostAuth(WorkerManager):
    """The worker is started once for the module, before any test's host exists: give it the hosts' (test) keys."""

    keep_alive = True   # each test's gateway shuts down (and would stop the worker); the module keeps it for all tests

    def _host_settings(self):
        return SimpleNamespace(enable_auth=True, api_key=ADMIN_KEY, api_keys=KEYS, api_key_header="X-API-Key")

    def stop(self):
        if not self.keep_alive:
            super().stop()


@pytest.fixture(scope="module")
def real_manager(tmp_path_factory):
    if not (VENDOR / "app" / "main.py").is_file():
        pytest.skip("the vendored CHARGE_SHEET_OCR snapshot is not present")
    patch = pytest.MonkeyPatch()
    for k, v in DETERMINISTIC.items():
        patch.setenv(k, v)
    data = tmp_path_factory.mktemp("real_worker")
    settings = load_settings({"UNIFIED_DATA_DIR": str(data), "CS_WORKER_URL": f"http://127.0.0.1:{free_port()}",
                              "CS_WORKER_PYTHON": worker_python(), "CS_WORKER_STARTUP_SECONDS": "120"})
    manager = HostAuth(settings)
    try:
        manager.ensure_running()
    except Exception as exc:  # noqa: BLE001
        manager.keep_alive = False
        manager.stop()
        patch.undo()
        pytest.skip(f"the charge-sheet worker could not be started with {worker_python()}: {exc}")
    yield manager
    manager.keep_alive = False
    manager.stop()
    patch.undo()


@pytest.fixture(scope="module")
def plain_cs(tmp_path_factory):
    """A plain CHARGE_SHEET_OCR application (no wrapper, no probe route) for the same-input comparison."""
    patch = pytest.MonkeyPatch()
    data = tmp_path_factory.mktemp("plain_cs")
    env = {k: v for k, v in os.environ.items() if k.upper() not in ("PYTHONPATH", "PYTHONHOME", "VIRTUAL_ENV")}
    env.update(DETERMINISTIC)
    env.update({"DATABASE_URL": "sqlite:///" + (data / "plain.db").as_posix(), "UPLOAD_DIR": str(data / "uploads"),
                "ENABLE_AUTH": "false", "PYTHONUTF8": "1"})
    (data / "uploads").mkdir()
    port = free_port()
    proc = subprocess.Popen([worker_python(), "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(port)],
                            cwd=str(VENDOR), env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    url = f"http://127.0.0.1:{port}"
    deadline = time.time() + 90
    while time.time() < deadline:
        try:
            if httpx.get(f"{url}/health", timeout=2).status_code == 200:
                break
        except httpx.HTTPError:
            time.sleep(0.5)
    else:
        proc.kill()
        pytest.skip("a plain CHARGE_SHEET_OCR application could not be started")
    yield url
    proc.terminate()
    proc.wait(timeout=15)
    patch.undo()


@pytest.fixture()
def real(make_client, out_root, real_manager, tmp_path):
    host = make_client(enable_auth=True).app
    settings = load_settings({"UNIFIED_DATA_DIR": str(tmp_path / "unified"), "CS_WORKER_URL": real_manager.url,
                              "CS_WORKER_MANAGE": "false", "CLASSIFY_OCR": "false"})
    gateway = create_gateway(settings, host_app=host, manager=real_manager)
    with TestClient(gateway) as client:
        client.headers.update({"X-API-Key": ADMIN_KEY})
        yield SimpleNamespace(client=client, gateway=gateway, settings=settings, manager=real_manager)
    gateway.coordinator.report.shutdown()


def sample(name: str) -> bytes:
    return (SAMPLES / name).read_bytes()


def settled(client, doc_id):
    """The native results once the worker has finished with the document."""
    def done():
        body = client.get(f"/api/v1/documents/{doc_id}").json()["document"]
        return body if body["status"] in ("completed", "failed", "verified") else None

    return wait_for(done, timeout=120)


# Values that legitimately differ between two runs of the same service: clocks, and database state
# (`previously_seen` / `template_id` say whether THAT service's database had already met the document's template,
# which depends on what was uploaded to it before - not on how the document was read).
VOLATILE = {"created_at", "updated_at", "processed_at", "elapsed_seconds", "duration_ms", "duration_seconds",
            "started_at", "finished_at", "elapsed_ms", "timestamp", "extracted_at", "previously_seen", "template_id"}


def normalise(node):
    node = copy.deepcopy(node)
    if isinstance(node, dict):
        return {k: normalise(v) for k, v in node.items() if k not in VOLATILE}
    if isinstance(node, list):
        return [normalise(v) for v in node]
    return node


# ------------------------------------------------------------------------------------------------ routing
def test_real_snapshot_routes_and_classifies_through_the_probe(real):
    r = upload(real.client, "sample_01_digital_standard.pdf", sample("sample_01_digital_standard.pdf"))
    assert r.status_code == 202, r.text
    body = r.json()
    assert body["document_type"] == "charge_sheet"
    detector = body["classification"]["signals"]["charge_sheet"]["detector"]
    assert detector is not None and detector["document_type"] == "charge_sheet"      # the snapshot's own DocumentDetector
    assert body["classification"]["text"]["source"] == "native"
    assert real.gateway.coordinator.registry.adapter_of(body["document_id"]) == "charge_sheet_ocr"


@pytest.mark.parametrize("name", ["sample_01_digital_standard.pdf", "sample_02_alternate_labels.pdf",
                                  "sample_03_multipage.pdf", "sample_05_header_table_narrative_refs.pdf",
                                  "sample_06_missing_fields_and_traps.pdf"])
def test_unified_result_equals_the_plain_application_result(real, plain_cs, name):
    """Same file, same service code: the unified path must not change a single extracted value or field."""
    data = sample(name)
    doc_id = upload(real.client, name, data, document_type="charge_sheet").json()["document_id"]
    settled(real.client, doc_id)
    unified_payload = real.client.get(f"/api/v1/documents/{doc_id}/results").json()

    plain = httpx.post(f"{plain_cs}/api/v1/documents/upload", files={"file": (name, data, "application/pdf")}, timeout=120)
    assert plain.status_code == 202, plain.text
    plain_id = plain.json()["document_id"]
    deadline = time.time() + 120
    while time.time() < deadline:
        status = httpx.get(f"{plain_cs}/api/v1/documents/{plain_id}").json()["document"]["status"]
        if status in ("completed", "failed"):
            break
        time.sleep(0.5)
    plain_payload = httpx.get(f"{plain_cs}/api/v1/documents/{plain_id}/results").json()

    assert unified_payload["results"]["fields"], "the sample produced no fields at all"
    assert normalise(unified_payload["results"]) == normalise(plain_payload["results"])
    assert unified_payload["results"]["schema_version"] == plain_payload["results"]["schema_version"]
    # ... and the combined export embeds exactly that payload
    rep = wait_report(real.client, doc_id, ocr_status=("succeeded" if unified_payload["results"]["extraction_status"] != "partial" else "partial"))
    assert normalise(rep["ocr"]["payload"]["results"]) == normalise(plain_payload["results"])
    assert envelope_errors(rep) == []


# ---------------------------------------------------------------------------- review through the real service
def test_corrections_verification_and_reviewer_attribution_with_the_real_service(real):
    name = "sample_01_digital_standard.pdf"
    doc_id = upload(real.client, name, sample(name)).json()["document_id"]
    settled(real.client, doc_id)
    fields = real.client.get(f"/api/v1/documents/{doc_id}/results").json()["results"]["fields"]
    target = next(f for f in fields if f.get("value"))
    new_value = "Corrected value\nwith a second line  and  double spaces"
    r = real.client.patch(f"/api/v1/documents/{doc_id}/fields", headers={"X-API-Key": OPERATOR_KEY},
                          json={"corrections": [{"field_id": target["id"], "corrected_value": new_value, "review_status": "corrected"}]})
    assert r.status_code == 200, r.text

    # an operator may not verify; the reviewer may, and the verification is attributed to the reviewer by name
    verify = {"verified": True, "notes": "checked against the original"}
    assert real.client.post(f"/api/v1/documents/{doc_id}/verify", json=verify, headers={"X-API-Key": OPERATOR_KEY}).status_code == 403
    ok = real.client.post(f"/api/v1/documents/{doc_id}/verify", json=verify, headers={"X-API-Key": REVIEWER_KEY})
    assert ok.status_code == 200, ok.text
    assert ok.json()["document"]["verified_by"] == "rita"

    history = real.client.get(f"/api/v1/documents/{doc_id}/corrections").json()
    assert any(h["corrected_value"] == new_value for h in history)                  # kept as typed, incl. newline and spacing
    assert any(h["canonical_name"] == "_verification" for h in history)

    rep = wait_for(lambda: (lambda d: d if d["review"]["status"] == "verified" else None)(report(real.client, doc_id)))
    assert rep["review"]["verified_by"] == "rita" and rep["review"]["corrected_fields"] == 1
    assert rep["ocr"]["payload"]["field_corrections"][target["id"]]["corrected_value"] == new_value
    # editing a verified result re-opens review (the service's own rule, visible through the gateway)
    real.client.patch(f"/api/v1/documents/{doc_id}/fields", headers={"X-API-Key": OPERATOR_KEY},
                      json={"corrections": [{"field_id": target["id"], "corrected_value": "again"}]})
    wait_for(lambda: report(real.client, doc_id)["review"]["verified"] is False)


def test_every_native_output_endpoint_is_reachable_through_the_gateway(real):
    name = "sample_03_multipage.pdf"
    data = sample(name)
    doc_id = upload(real.client, name, data).json()["document_id"]
    settled(real.client, doc_id)
    base = f"/api/v1/documents/{doc_id}"
    fields_json = real.client.get(f"{base}/json")
    assert fields_json.status_code == 200 and isinstance(fields_json.json(), dict)
    assert real.client.get(f"{base}/json?download=true").headers["content-disposition"].endswith('.json"')
    assert real.client.get(f"{base}/raw").status_code == 200
    assert real.client.get(f"{base}/structured").status_code == 404                  # VLM off: the service's own answer
    source = real.client.get(f"{base}/source")
    assert source.status_code == 200 and source.content == data                       # the original, byte for byte
    page = real.client.get(f"{base}/pages/1/image?width=400")
    assert page.status_code == 200 and page.content[:4] == b"\x89PNG"
    assert real.client.get(f"{base}/pages/999/image").status_code == 404
    assert real.client.get(f"{base}/corrections").json() == []
    page_html = real.client.get(f"/results/{doc_id}")
    assert page_html.status_code == 200 and '"/cs-static/js/app.js' in page_html.text
    assert real.client.get("/cs-static/js/app.js").status_code == 200
    assert real.client.get("/cs-static/css/styles.css").status_code == 200


# ------------------------------------------------------------------------------------------- file types
def charge_sheet_docx(tmp_path) -> bytes:
    """A DOCX with enough text for the service to read (it refuses, rather than invents, near-empty documents)."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("CHARGE SHEET (Under Section 173 Cr.P.C.)")
    doc.add_paragraph("In the Court of the Hon'ble Judicial First Class Magistrate, Medak. " * 3)
    table = doc.add_table(rows=3, cols=2)
    for row, (label, value) in zip(table.rows, [("District", "Medak"), ("Police Station", "Havelighanapur"),
                                                 ("FIR No.", "185/2021")]):
        row.cells[0].text, row.cells[1].text = label, value
    doc.add_paragraph("Particulars of the accused: A1 Ravi Kumar, S/o Ramesh, aged 34 years, resident of Medak. " * 3)
    path = tmp_path / "cs.docx"
    doc.save(path)
    return path.read_bytes()


def test_docx_is_accepted_by_the_charge_sheet_service(real, tmp_path):
    r = upload(real.client, "converted.docx", charge_sheet_docx(tmp_path))
    assert r.json()["document_type"] == "charge_sheet"                                # recognised from the DOCX text
    doc_id = r.json()["document_id"]
    assert settled(real.client, doc_id)["status"] == "completed"
    assert real.client.get(f"/api/v1/documents/{doc_id}/results").json()["results"]["fields"]


def test_a_near_empty_docx_is_refused_by_the_service_and_reported_as_failed(real, tmp_path):
    """The service's own rule (no invented values): the unified export reports the failure, with no payload."""
    doc_id = upload(real.client, "tiny.docx", ir_docx(tmp_path), document_type="charge_sheet").json()["document_id"]
    assert settled(real.client, doc_id)["status"] == "failed"
    rep = wait_report(real.client, doc_id, ocr_status="failed")
    assert rep["ocr"]["payload"] is None and rep["processing"]["ocr_error_summary"]


def test_scanned_image_without_ocr_is_reported_as_it_is_never_as_success(real):
    doc_id = upload(real.client, "page.png", make_png((250, 250, 250), size=(600, 800)), document_type="charge_sheet").json()["document_id"]
    native = settled(real.client, doc_id)
    rep = wait_report(real.client, doc_id, image_extraction_status="not_applicable")
    # whatever the service decided (failed or partial), the export says the same; charge sheets have no image pipeline
    assert native["status"] in ("completed", "failed")
    if native["status"] == "failed":
        assert rep["processing"]["ocr_status"] == "failed" and rep["ocr"]["payload"] is None
    else:
        assert rep["processing"]["ocr_status"] in ("partial", "failed")


def test_an_interrogation_report_does_not_touch_the_charge_sheet_service(real, tmp_path, fake_faces):
    def total():
        return httpx.get(f"{real.manager.url}/api/v1/documents?limit=1", headers={"X-API-Key": ADMIN_KEY}, timeout=10).json()["total"]

    before = total()
    r = upload(real.client, "ir.docx", ir_docx(tmp_path))
    assert r.json()["document_type"] == "interrogation_report"
    wait_report(real.client, r.json()["document_id"], ocr_status="succeeded")
    assert total() == before
