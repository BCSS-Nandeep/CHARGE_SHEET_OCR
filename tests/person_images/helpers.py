"""Helpers for the person-image integration tests: synthetic documents, a stand-in face decision, polling.

Documents are synthetic. Image content is coloured noise; for the pipeline tests the *face decision* is replaced by a
deterministic stand-in (red-dominant image = "person"), so extraction, filtering, storage, statuses, JSON and UI are
exercised for real while the YuNet model itself is checked separately (negative control, and an optional real-face
image through the IR_TEST_FACE_IMAGE environment variable).
"""

from __future__ import annotations

import io
import time
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("cv2", reason="opencv-python-headless is required for the person-image integration")

import numpy as np  # noqa: E402
import pymupdf  # noqa: E402
from PIL import Image  # noqa: E402

from integrations.person_images import extractor, service, store  # noqa: E402
from integrations.person_images.wiring import attach  # noqa: E402
from tests.conftest import ADMIN_KEY  # noqa: E402

RED, BLUE = (200, 30, 30), (30, 30, 200)


def make_png(color, size=(160, 200), seed=0) -> bytes:
    rng = np.random.default_rng(seed)
    arr = np.clip(np.array(color, dtype=np.int16) + rng.integers(-10, 10, size=(size[1], size[0], 3)), 0, 255).astype("uint8")
    buf = io.BytesIO()
    Image.fromarray(arr).save(buf, "PNG")
    return buf.getvalue()


def fake_detect(blob: bytes):
    """Stand-in for the face decision: a red-dominant image 'contains a person'. Size rules stay the real ones."""
    with Image.open(io.BytesIO(blob)) as im:
        rgb = np.array(im.convert("RGB"))
    h, w = rgb.shape[:2]
    if h < extractor.MIN_IMAGE_SIZE or w < extractor.MIN_IMAGE_SIZE:
        return w, h, 0
    r, g, b = rgb[..., 0].mean(), rgb[..., 1].mean(), rgb[..., 2].mean()
    return w, h, 1 if (r > 120 and g < 90 and b < 90) else 0


def make_docx(path: Path, images=()):
    """A report-like DOCX (headings + a table) with the given PNG images embedded."""
    from docx import Document
    from docx.shared import Cm

    doc = Document()
    doc.add_paragraph("Annexure-I")
    doc.add_paragraph("Interrogation Report")
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Name", "Test Person"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Age", "30"
    for i, png in enumerate(images):
        p = io.BytesIO(png)
        doc.add_picture(p, width=Cm(3 + i * 0.1))
    doc.save(path)
    return path


def make_pdf(path: Path, images=(), text="Interrogation Report " * 20):
    doc = pymupdf.open()
    page = doc.new_page()
    if text:
        page.insert_text((72, 72), text)
    y = 120
    for png in images:
        page.insert_image(pymupdf.Rect(72, y, 72 + 80, y + 100), stream=png)
        y += 110
    doc.save(path)
    doc.close()
    return path


def make_xlsx_bytes(images) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        for i, png in enumerate(images, start=1):
            z.writestr(f"xl/media/image{i}.png", png)
    return buf.getvalue()


def upload(client, path: Path, content_type=None):
    ct = content_type or ("application/pdf" if path.suffix == ".pdf" else
                          "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    r = client.post("/api/v1/documents/upload", files={"file": (path.name, path.read_bytes(), ct)})
    return r


def wait_images(client, doc_id, timeout=30.0):
    """Poll the new status endpoint until the image run is no longer queued / processing."""
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        r = client.get(f"/api/v1/documents/{doc_id}/person-images")
        if r.status_code == 200 and r.json()["status"] not in ("queued", "processing", "not_started"):
            service.wait_idle()
            return r.json()
        time.sleep(0.1)
    raise AssertionError(f"image extraction did not finish: {r.status_code} {r.text[:200]}")


def wait_for(predicate, timeout=10.0, what="condition"):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        value = predicate()
        if value:
            return value
        time.sleep(0.1)
    raise AssertionError(f"timed out waiting for {what}")


def report_file(out_root: Path, doc_id: str) -> Path:
    """Path of the combined JSON in the document's folder (named after the uploaded file)."""
    folder = store.find_dir(doc_id)
    return (folder if folder is not None else out_root / "_none_") / "report_result.json"


def doc_folder(doc_id: str) -> Path:
    folder = store.find_dir(doc_id)
    assert folder is not None, "the document has no output folder"
    return folder