"""End-to-end through the API with a fake model server: structured results, scanned pages read
by the VLM, partial / failed runs, corrections, and image / DOCX uploads."""

from __future__ import annotations

import io
from pathlib import Path

import fitz
import pytest
from fastapi.testclient import TestClient
from PIL import Image

from tests.conftest import ADMIN_KEY, OPERATOR_KEY, REVIEWER_KEY
from tests.fake_vlm import FakeVlm

NANONETS_S04 = (  # what the real model returned for the synthetic scanned fixture's page 1
    "CHARGE SHEET\nDistrict: Coastal District\nPolice Station: Harbour PS\nFIR No.: 901/2021\nYear: 2021\n"
    "Charge Sheet No.: CS-3/2022\nDate of FIR: 18/08/2021\nCourt: Special Court\nu/s 307 IPC\n"
    "Complainant: Arjun Nair\nInvestigating Officer: DySP L. Thomas\nParticulars of the Accused:\n"
    "1. Manoj Pillai s/o Pillai age 35\nList of Witnesses:\n1. Arjun Nair\n2. Constable Ravi\n"
    "Brief Facts: Assault case investigated and charge sheeted.\nConclusion: Accused sent for trial."
)


@pytest.fixture()
def vlm(monkeypatch):
    """Enable the VLM path and replace the client with a fake; returns a configurator."""
    monkeypatch.setenv("VLM_ENABLED", "true")
    monkeypatch.setenv("VLM_BASE_URL", "http://127.0.0.1:8081/v1")
    monkeypatch.setenv("VLM_MODEL", "fake-model")
    from app.config import get_settings

    get_settings.cache_clear()

    def install(fake: FakeVlm) -> FakeVlm:
        monkeypatch.setattr("app.services.extraction_service.VlmClient", fake)
        monkeypatch.setattr("app.api.routes_health.VlmClient", fake)
        return fake

    return install


def upload(client: TestClient, name: str, data: bytes, ctype: str, key: str = ADMIN_KEY):
    res = client.post("/api/v1/documents/upload", files={"file": (name, data, ctype)}, headers={"X-API-Key": key})
    return res


def upload_path(client, path: Path, key: str = ADMIN_KEY) -> str:
    res = upload(client, path.name, path.read_bytes(), "application/pdf", key)
    assert res.status_code == 202, res.text
    return res.json()["document_id"]


def results(client, doc_id):
    res = client.get(f"/api/v1/documents/{doc_id}/results")
    assert res.status_code == 200, res.text
    return res.json()


def sv(body, path):
    sec, *rest = path.split(".")
    s = body["results"]["structured"]
    if len(rest) == 1:
        return s["single"][sec][rest[0]]
    return s["lists"][sec][int(rest[0])]["fields"][rest[1]]


# ---------------------------------------------------------------- digital PDF


def test_digital_pdf_gets_structured_fields_with_evidence(client, fixtures_dir, vlm):
    fake = vlm(FakeVlm())
    body = results(client, upload_path(client, fixtures_dir / "sample_01_digital_standard.pdf"))
    r = body["results"]
    assert r["fields"], "as-printed fields are still produced"
    s = r["structured"]
    assert s["schema_version"] == "1.0" and s["document_status"] == "extracted"
    fir = sv(body, "case.fir_number")
    assert fir["value"] == "123/2024" and fir["status"] == "extracted" and fir["source_page"] == 1
    assert "123/2024" in fir["source_text"] and fir["text_engine"] == "pdf_text"
    assert sv(body, "case.crime_number")["status"] == "not_found"  # never stated: not invented
    accused = s["lists"]["accused"]
    assert [a["fields"]["name"]["value"] for a in accused] == ["Suresh Singh", "Vijay Patel"]
    assert accused[1]["fields"]["age"]["value"] == "28"
    assert len(fake.map_calls) == 2  # every page is mapped, one request each
    assert body["structured_layout"][0]["key"] == "case"
    assert any(st["stage"] == "structured_mapping" and st["status"] == "ok" for st in r["stage_results"])


def test_narrative_fields_are_linked_not_copied(client, fixtures_dir, vlm):
    vlm(FakeVlm())
    body = results(client, upload_path(client, fixtures_dir / "sample_01_digital_standard.pdf"))
    ids = body["results"]["structured"]["narrative_field_ids"]
    fields = {f["id"]: f for f in body["results"]["fields"]}
    assert ids and all(fields[i]["text_block"] for i in ids)


def test_every_document_has_the_same_structured_sections(client, fixtures_dir, vlm):
    vlm(FakeVlm())
    a = results(client, upload_path(client, fixtures_dir / "sample_01_digital_standard.pdf"))["results"]["structured"]
    b = results(client, upload_path(client, fixtures_dir / "sample_03_multipage.pdf"))["results"]["structured"]
    assert {k: sorted(v) for k, v in a["single"].items()} == {k: sorted(v) for k, v in b["single"].items()}
    assert list(a["lists"]) == list(b["lists"])


def test_multipage_records_keep_their_source_pages(client, fixtures_dir, vlm):
    vlm(FakeVlm())
    body = results(client, upload_path(client, fixtures_dir / "sample_03_multipage.pdf"))
    s = body["results"]["structured"]
    assert sv(body, "case.fir_number")["source_page"] == 1
    assert [a["pages"] for a in s["lists"]["accused"]] == [[2], [2], [2]]
    assert s["lists"]["witnesses"][0]["pages"] == [3]
    assert s["lists"]["acts_and_sections"][0]["pages"] == [3]
    assert sv(body, "case.fir_date")["status"] == "not_found"


def test_without_the_vlm_nothing_changes(client, fixtures_dir):
    body = results(client, upload_path(client, fixtures_dir / "sample_01_digital_standard.pdf"))
    assert body["results"]["structured"] is None and body["results"]["fields"]


# ---------------------------------------------------------------- scanned PDF read by the VLM


def test_scanned_pdf_is_read_by_the_vlm_and_mapped(client, fixtures_dir, vlm):
    fake = vlm(FakeVlm(transcripts={0: NANONETS_S04}))
    body = results(client, upload_path(client, fixtures_dir / "sample_04_scanned_image.pdf"))
    r = body["results"]
    page = r["page_summaries"][0]
    assert page["text_source"] == "ocr" and page["ocr_engine"] == "nanonets" and page["ocr_confidence"] is None
    assert fake.transcribe_calls == 1
    assert sv(body, "case.fir_number")["value"] == "901/2021"
    assert sv(body, "case.fir_number")["text_engine"] == "nanonets"
    assert r["structured"]["lists"]["accused"][0]["fields"]["name"]["value"] == "Manoj Pillai"
    assert any(f["label"] and "District" in f["label"] for f in r["fields"])  # as-printed fields from the VLM text


def test_raw_json_keeps_the_vlm_lines_as_read(client, fixtures_dir, vlm):
    vlm(FakeVlm(transcripts={0: "<watermark>COPY</watermark>\n" + NANONETS_S04}))
    doc_id = upload_path(client, fixtures_dir / "sample_04_scanned_image.pdf")
    raw = client.get(f"/api/v1/documents/{doc_id}/raw").json()
    lines = raw["pages"][0]["lines"]
    assert {ln["source"] for ln in lines} == {"vlm"} and raw["pages"][0]["ocr_engine"] == "nanonets"
    assert lines[0]["text"] == "<watermark>COPY</watermark>"  # tags kept in the raw view


def test_vlm_failure_on_a_scanned_page_is_a_clear_failure(client, fixtures_dir, vlm):
    vlm(FakeVlm(fail_transcribe_pages={1}))
    doc_id = upload_path(client, fixtures_dir / "sample_04_scanned_image.pdf")
    res = client.get(f"/api/v1/documents/{doc_id}/results")
    assert res.status_code == 422
    msg = res.json()["detail"]["message"]
    assert "vlm_timeout" in msg and "page 1" in msg
    assert client.get(f"/api/v1/documents/{doc_id}").json()["document"]["status"] == "failed"


def test_unavailable_server_fails_scanned_pages_without_pretending(client, fixtures_dir, vlm):
    fake = vlm(FakeVlm(unavailable=True))
    doc_id = upload_path(client, fixtures_dir / "sample_04_scanned_image.pdf")
    res = client.get(f"/api/v1/documents/{doc_id}/results")
    assert res.status_code == 422 and fake.transcribe_calls == 0  # nothing was sent
    assert "not reachable" in res.json()["detail"]["message"]


def test_unreadable_blank_vlm_page_is_failed_not_empty_success(client, fixtures_dir, vlm):
    vlm(FakeVlm(transcripts={0: "   "}))
    doc_id = upload_path(client, fixtures_dir / "sample_04_scanned_image.pdf")
    assert client.get(f"/api/v1/documents/{doc_id}/results").status_code == 422


def test_fallback_engine_is_used_only_when_configured(monkeypatch, fixtures_dir, vlm):
    """With VLM_FALLBACK_ENGINE=paddleocr a failed VLM page goes to the local OCR (and says so)."""
    from app.config import get_settings
    from app.services.extraction_service import ExtractionService
    from app.services.pdf_service import PageContent
    from app.schemas.charge_sheet import PageSummary

    fake = FakeVlm(unavailable=True)
    vlm(fake)

    def acquire(fallback: str):
        monkeypatch.setenv("VLM_FALLBACK_ENGINE", fallback)
        get_settings.cache_clear()
        svc = ExtractionService()
        svc._vlm_down = "server not reachable"
        page = PageContent(page_number=1, text="", width=1, height=1, rotation=0, char_count=0, route="ocr")
        summary = PageSummary(page_number=1, route="ocr")
        summaries, failed = [], []
        handled = svc._vlm_acquire_page(fixtures_dir / "sample_04_scanned_image.pdf", page, summary, summaries, {}, failed, 1, 1)
        return handled, summary, failed

    handled, summary, failed = acquire("none")
    assert handled and summary.page_status == "failed" and failed and "unavailable" in summary.note
    handled, summary, failed = acquire("paddleocr")
    assert not handled and not failed and "PaddleOCR used" in summary.note


# ---------------------------------------------------------------- partial results


def test_one_failed_page_keeps_the_others_and_marks_the_document_partial(client, fixtures_dir, vlm):
    vlm(FakeVlm(fail_map_pages={2}))
    body = results(client, upload_path(client, fixtures_dir / "sample_03_multipage.pdf"))
    r = body["results"]
    s = r["structured"]
    assert r["extraction_status"] == "partial" and s["document_status"] == "partial"
    assert sv(body, "case.fir_number")["value"] == "55/2022"  # page 1 kept
    assert s["lists"]["witnesses"][0]["pages"] == [3]  # page 3 kept
    assert s["lists"]["accused"] == []  # page 2's accused are missing, not made up
    assert [p["mapping_status"] for p in s["pages"]] == ["ok", "failed", "ok"]
    assert any(i["code"] == "page_failed" and i["page"] == 2 for i in s["validation"]["issues"])
    assert any(w.startswith("STRUCTURED_MAPPING_PARTIAL") for w in r["extraction_warnings"])


def test_unavailable_server_on_a_digital_pdf_gives_fields_but_no_structured_values(client, fixtures_dir, vlm):
    fake = vlm(FakeVlm(unavailable=True))
    body = results(client, upload_path(client, fixtures_dir / "sample_01_digital_standard.pdf"))
    r = body["results"]
    assert r["fields"] and r["extraction_status"] == "partial"
    assert r["structured"]["document_status"] == "failed"
    assert all(v["status"] == "not_found" for sec in r["structured"]["single"].values() for v in sec.values())
    assert any(w.startswith("VLM_UNAVAILABLE") for w in r["extraction_warnings"])
    assert fake.map_calls == []


def test_server_that_stops_answering_is_not_asked_for_every_page(client, fixtures_dir, vlm, monkeypatch):
    monkeypatch.setenv("VLM_ABORT_AFTER_FAILURES", "2")
    fake = vlm(FakeVlm(fail_map_pages={1, 2, 3}))
    results(client, upload_path(client, fixtures_dir / "sample_03_multipage.pdf"))
    assert len(fake.map_calls) == 2  # third page was skipped after two consecutive timeouts


# ---------------------------------------------------------------- corrections and review


def test_structured_value_correction_is_kept_apart_from_the_machine_value(client, fixtures_dir, vlm):
    vlm(FakeVlm())
    doc_id = upload_path(client, fixtures_dir / "sample_01_digital_standard.pdf")
    res = client.patch(f"/api/v1/documents/{doc_id}/fields", headers={"X-API-Key": OPERATOR_KEY}, json={
        "corrections": [
            {"field_id": "s.case.fir_number", "corrected_value": "124/2024"},
            {"field_id": "s.case.crime_number", "corrected_value": "77/2024"},  # was not_found
            {"field_id": "s.accused.1.age", "corrected_value": "29"},
        ]})
    assert res.status_code == 200, res.text
    body = res.json()
    fir = sv(body, "case.fir_number")
    assert fir["value"] == "123/2024" and fir["corrected_value"] == "124/2024" and fir["review_status"] == "corrected"
    assert fir["status"] == "extracted"  # the machine status is not rewritten
    crime = sv(body, "case.crime_number")
    assert crime["value"] is None and crime["corrected_value"] == "77/2024"
    assert sv(body, "accused.1.age")["corrected_value"] == "29"
    hist = client.get(f"/api/v1/documents/{doc_id}/corrections").json()
    assert {h["canonical_name"]: h["original_value"] for h in hist}["s.case.fir_number"] == "123/2024"
    assert {h["canonical_name"]: h["original_value"] for h in hist}["s.case.crime_number"] == ""


def test_unknown_structured_path_is_rejected(client, fixtures_dir, vlm):
    vlm(FakeVlm())
    doc_id = upload_path(client, fixtures_dir / "sample_01_digital_standard.pdf")
    for bad in ("s.case.nope", "s.accused.9.name", "s.nothing.x"):
        res = client.patch(f"/api/v1/documents/{doc_id}/fields", json={
            "corrections": [{"field_id": bad, "corrected_value": "x"}]})
        assert res.status_code == 400, bad


def test_corrections_survive_reprocess_only_when_the_machine_value_is_unchanged(client, fixtures_dir, vlm):
    vlm(FakeVlm())
    doc_id = upload_path(client, fixtures_dir / "sample_01_digital_standard.pdf")
    client.patch(f"/api/v1/documents/{doc_id}/fields", json={"corrections": [
        {"field_id": "s.case.fir_number", "corrected_value": "124/2024"}]})
    client.post(f"/api/v1/documents/{doc_id}/reprocess")
    assert sv(results(client, doc_id), "case.fir_number")["corrected_value"] == "124/2024"
    # the model now reads a different FIR number: the old correction no longer applies
    vlm(FakeVlm(answer=lambda t: {"case": {"fir_number": "999/2024"}}))
    client.post(f"/api/v1/documents/{doc_id}/reprocess")
    body = results(client, doc_id)
    assert sv(body, "case.fir_number")["corrected_value"] is None
    assert any("not re-applied" in w for w in body["results"]["extraction_warnings"])


def test_verification_still_needs_a_reviewer_and_notes(client, fixtures_dir, vlm):
    vlm(FakeVlm())
    doc_id = upload_path(client, fixtures_dir / "sample_01_digital_standard.pdf")
    assert client.post(f"/api/v1/documents/{doc_id}/verify", json={"verified": True, "notes": "ok"},
                       headers={"X-API-Key": OPERATOR_KEY}).status_code == 403
    assert client.post(f"/api/v1/documents/{doc_id}/verify", json={"verified": True},
                       headers={"X-API-Key": REVIEWER_KEY}).status_code == 400
    # a fresh extraction is never presented as verified
    body = results(client, doc_id)
    assert body["document"]["verified"] is False
    assert all(v["review_status"] != "verified" for sec in body["results"]["structured"]["single"].values()
               for v in sec.values())


# ---------------------------------------------------------------- the one final JSON


def test_final_structured_json_is_one_document_with_final_values_and_evidence(client, fixtures_dir, vlm):
    vlm(FakeVlm())
    doc_id = upload_path(client, fixtures_dir / "sample_01_digital_standard.pdf")
    client.patch(f"/api/v1/documents/{doc_id}/fields", json={"corrections": [
        {"field_id": "s.case.fir_number", "corrected_value": "124/2024"},
        {"field_id": "s.case.crime_number", "corrected_value": "77/2024"}]})
    res = client.get(f"/api/v1/documents/{doc_id}/structured")
    assert res.status_code == 200
    j = res.json()
    assert j["document_id"] == doc_id and j["schema_version"] == "1.0" and j["verification"]["verified"] is False
    fir = j["case"]["fir_number"]
    assert (fir["value"], fir["machine_value"], fir["human_corrected"]) == ("124/2024", "123/2024", True)
    assert fir["status"] == "extracted" and fir["source_page"] == 1 and "123/2024" in fir["source_text"]
    assert j["case"]["crime_number"]["value"] == "77/2024" and j["case"]["crime_number"]["machine_value"] is None
    assert j["case"]["district"]["value"] == "Sample District" and not j["case"]["district"]["human_corrected"]
    assert [a["fields"]["name"]["value"] for a in j["accused"]] == ["Suresh Singh", "Vijay Patel"]
    assert set(j) >= {"case", "complainant", "informant", "acts_and_sections", "investigating_officers", "accused",
                      "witnesses", "property", "other_details", "pages", "validation"}
    dl = client.get(f"/api/v1/documents/{doc_id}/structured?download=true")
    assert "attachment" in dl.headers["content-disposition"] and dl.json() == j


def test_final_structured_json_is_unavailable_without_a_structured_result(client, fixtures_dir):
    doc_id = upload_path(client, fixtures_dir / "sample_01_digital_standard.pdf")  # VLM off
    res = client.get(f"/api/v1/documents/{doc_id}/structured")
    assert res.status_code == 404 and "VLM_ENABLED" in res.json()["detail"]
    assert client.get(f"/api/v1/documents/{doc_id}/structured", headers={"X-API-Key": ""}).status_code == 401


# ---------------------------------------------------------------- health


def test_health_reports_the_model(client, vlm):
    vlm(FakeVlm())
    h = client.get("/health").json()
    assert h["vlm_enabled"] and h["vlm_available"] and h["vlm_model"] == "fake-model"
    vlm(FakeVlm(unavailable=True))
    h = client.get("/health").json()
    assert h["vlm_enabled"] and not h["vlm_available"] and "not reachable" in h["vlm_error"]


# ---------------------------------------------------------------- image uploads


def png_bytes(w=400, h=300, color=(255, 255, 255)):
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, "PNG")
    return buf.getvalue()


def test_png_upload_is_stored_untouched_converted_and_read_by_the_vlm(client, vlm):
    fake = vlm(FakeVlm(transcripts={0: NANONETS_S04}))
    data = png_bytes()
    res = upload(client, "scan.png", data, "image/png")
    assert res.status_code == 202
    doc_id = res.json()["document_id"]
    meta = client.get(f"/api/v1/documents/{doc_id}").json()["document"]
    assert meta["stored_filename"].endswith(".png") and meta["content_type"] == "image/png"
    src = client.get(f"/api/v1/documents/{doc_id}/source")
    assert src.status_code == 200 and src.headers["content-type"] == "image/png" and src.content == data
    body = results(client, doc_id)
    assert body["results"]["page_summaries"][0]["ocr_engine"] == "nanonets"
    assert sv(body, "case.fir_number")["value"] == "901/2021"
    assert any(w.startswith("IMAGE_SOURCE") for w in body["results"]["extraction_warnings"])
    assert fake.transcribe_calls == 1


def test_multiframe_tiff_becomes_one_page_per_frame(client, vlm):
    fake = vlm(FakeVlm(transcripts={0: NANONETS_S04}))
    frames = [Image.new("RGB", (300, 400), c) for c in ((255, 255, 255), (250, 250, 250), (245, 245, 245))]
    buf = io.BytesIO()
    frames[0].save(buf, "TIFF", save_all=True, append_images=frames[1:])
    res = upload(client, "multi.tiff", buf.getvalue(), "image/tiff")
    assert res.status_code == 202
    body = results(client, res.json()["document_id"])
    assert [p["page_number"] for p in body["results"]["page_summaries"]] == [1, 2, 3]
    assert fake.transcribe_calls == 3


def test_exif_rotated_photo_is_straightened(client, vlm, tmp_path):
    vlm(FakeVlm(transcripts={0: NANONETS_S04}))
    img = Image.new("RGB", (600, 200), (255, 255, 255))
    exif = Image.Exif()
    exif[0x0112] = 6  # rotate 90° clockwise to display
    buf = io.BytesIO()
    img.save(buf, "JPEG", exif=exif)
    res = upload(client, "photo.jpg", buf.getvalue(), "image/jpeg")
    doc_id = res.json()["document_id"]
    assert client.get(f"/api/v1/documents/{doc_id}/results").status_code == 200
    page = client.get(f"/api/v1/documents/{doc_id}/pages/1/image?width=300")
    assert page.status_code == 200 and page.headers["content-type"] == "image/png"
    out = Image.open(io.BytesIO(page.content))
    assert out.height > out.width  # landscape pixels, displayed portrait


def test_corrupt_image_fails_with_a_clear_message(client, vlm):
    vlm(FakeVlm())
    res = upload(client, "bad.png", b"\x89PNG\r\n\x1a\n" + b"garbage" * 20, "image/png")
    assert res.status_code == 202
    doc_id = res.json()["document_id"]
    doc = client.get(f"/api/v1/documents/{doc_id}").json()["document"]
    assert doc["status"] == "failed" and doc["error_message"]
    assert client.get(f"/api/v1/documents/{doc_id}/results").status_code == 422


@pytest.mark.parametrize("name,data", [
    ("notes.txt", b"just text"),
    ("fake.docx", b"PK\x03\x04 not really a zip"),
    ("archive.zip", b"PK\x03\x04" + b"\0" * 40),
    ("empty.png", b""),
    ("renamed.pdf", b"\x89PNG-not-really"),
])
def test_unsupported_or_malformed_uploads_are_rejected(client, name, data):
    res = upload(client, name, data, "application/octet-stream")
    assert res.status_code == 400
    assert "PDF, PNG, JPEG, TIFF or DOCX" in res.text or "empty" in res.text


def test_macos_shadow_file_gets_a_specific_message(client):
    appledouble = b"\x00\x05\x16\x07\x00\x02\x00\x00Mac OS X        " + b"\x00" * 140  # 163-byte "._name.pdf"
    res = upload(client, "._2022-04-26_scan.pdf", appledouble, "application/pdf")
    assert res.status_code == 400 and "'._'" in res.json()["detail"] and "without the" in res.json()["detail"]


def test_type_is_decided_by_content_not_by_name(client, vlm):
    vlm(FakeVlm(transcripts={0: NANONETS_S04}))
    res = upload(client, "actually_a_png.pdf", png_bytes(), "application/pdf")
    assert res.status_code == 202
    meta = client.get(f"/api/v1/documents/{res.json()['document_id']}").json()["document"]
    assert meta["stored_filename"].endswith(".png") and meta["content_type"] == "image/png"


# ---------------------------------------------------------------- DOCX


def make_docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("CHARGE SHEET / FINAL REPORT", level=1)
    table = doc.add_table(rows=4, cols=2)
    for row, (a, b) in zip(table.rows, [("District", "Sample District"), ("Police Station", "Central PS"),
                                        ("FIR No", "123/2024"), ("Year", "2024")]):
        row.cells[0].text, row.cells[1].text = a, b
    doc.add_paragraph("Particulars of the Accused:")
    doc.add_paragraph("1. Suresh Singh s/o Mohan Singh age 32")
    img = io.BytesIO()
    Image.new("RGB", (120, 60), (200, 30, 30)).save(img, "PNG")
    img.seek(0)
    doc.add_paragraph("Seizure photograph follows:")
    doc.add_picture(img)
    doc.add_paragraph("Brief Facts: an assault was reported.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_docx_keeps_tables_paragraphs_and_pictures_in_reading_order(client, vlm):
    vlm(FakeVlm(transcripts={0: ""}))
    data = make_docx()
    res = upload(client, "report.docx", data, "application/octet-stream")
    assert res.status_code == 202, res.text
    doc_id = res.json()["document_id"]
    src = client.get(f"/api/v1/documents/{doc_id}/source")
    assert src.content == data and "attachment" in src.headers["content-disposition"]
    body = results(client, doc_id)
    r = body["results"]
    assert any(w.startswith("DOCX_RENDITION") for w in r["extraction_warnings"])
    assert sv(body, "case.police_station")["value"] == "Central PS"
    assert sv(body, "case.fir_number")["value"] == "123/2024"
    assert r["structured"]["lists"]["accused"][0]["fields"]["name"]["value"] == "Suresh Singh"
    raw = client.get(f"/api/v1/documents/{doc_id}/raw").json()
    text = "\n".join(l["text"] for p in raw["pages"] for l in p["lines"])
    assert text.index("District") < text.index("Particulars of the Accused") < text.index("Seizure photograph") < text.index("Brief Facts")
    # the embedded picture is on the page the preview renders
    page = client.get(f"/api/v1/documents/{doc_id}/pages/1/image")
    assert page.status_code == 200 and page.content[:4] == b"\x89PNG"


def test_docx_picture_is_in_the_derived_pdf(tmp_path):
    from app.services.source_prep import docx_to_pdf

    src = tmp_path / "a.docx"
    src.write_bytes(make_docx())
    out = tmp_path / "a.pdf"
    assert docx_to_pdf(src, out) >= 1
    doc = fitz.open(out)
    assert sum(len(p.get_images()) for p in doc) == 1
    assert "123/2024" in doc[0].get_text()


def test_docx_with_only_unreadable_content_fails_cleanly(client, vlm):
    from docx import Document

    vlm(FakeVlm())
    buf = io.BytesIO()
    Document().save(buf)  # an empty Word document
    res = upload(client, "empty.docx", buf.getvalue(), "application/octet-stream")
    assert res.status_code == 202
    doc_id = res.json()["document_id"]
    assert client.get(f"/api/v1/documents/{doc_id}").json()["document"]["status"] == "failed"


# ---------------------------------------------------------------- page images


def test_page_image_endpoint(client, fixtures_dir):
    doc_id = upload_path(client, fixtures_dir / "sample_03_multipage.pdf")
    ok = client.get(f"/api/v1/documents/{doc_id}/pages/2/image?width=500")
    assert ok.status_code == 200 and ok.content[:8] == b"\x89PNG\r\n\x1a\n"
    assert client.get(f"/api/v1/documents/{doc_id}/pages/9/image").status_code == 404
    assert client.get(f"/api/v1/documents/{doc_id}/pages/1/image", headers={"X-API-Key": ""}).status_code == 401


def test_delete_removes_the_derived_pdf_too(client, vlm, tmp_path):
    vlm(FakeVlm(transcripts={0: NANONETS_S04}))
    res = upload(client, "scan.png", png_bytes(), "image/png")
    doc_id = res.json()["document_id"]
    from app.config import get_settings

    up = get_settings().upload_dir
    assert list(up.glob("*.png")) and list(up.glob("*.pages.pdf"))
    assert client.delete(f"/api/v1/documents/{doc_id}").status_code == 204
    assert not list(up.glob("*.png")) and not list(up.glob("*.pages.pdf"))
