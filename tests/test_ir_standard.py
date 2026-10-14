"""Standard-template IR: schema, validation, DOCX round trip, VLM mapping, API flow.

The VLM is replaced by an in-process fake OpenAI-compatible server
(httpx.MockTransport); no model is needed to run these tests.
"""

from __future__ import annotations

import io
import json
import re
from pathlib import Path

import httpx
import pytest

from app.config import get_settings
from app.ir_standard.docx_markdown import docx_to_markdown
from app.ir_standard.docx_writer import render_docx
from app.ir_standard.mapper import extract_json, map_markdown, parse_narrative
from app.ir_standard.paths import apply_overrides, flatten, set_value
from app.ir_standard.schema import FAMILY_RELATIONS, SECTION_BY_NUMBER, SECTIONS
from app.ir_standard.sectioner import split_sections
from app.ir_standard.validator import blank_standard, build_standard, snap_option
from app.ir_standard.vlm_client import VlmClient, VlmError


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def _vlm_env(monkeypatch):
    monkeypatch.setenv("LLM_ENABLED", "true")
    monkeypatch.setenv("LLM_BASE_URL", "http://127.0.0.1:9/v1")
    monkeypatch.setenv("LLM_MODEL", "fake-vlm")
    monkeypatch.setenv("IR_STANDARD_ENABLED", "true")
    get_settings.cache_clear()


def fake_server(answers, page_markdown=None, calls=None):
    """OpenAI-compatible fake. ``answers``: {section number: raw model answer (dict or str)}."""

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/models"):
            return httpx.Response(200, json={"data": []})
        body = json.loads(request.content)
        content = body["messages"][-1]["content"]
        if calls is not None:
            calls.append(content if isinstance(content, str) else "image")
        if isinstance(content, list):  # page image -> transcription
            text = (page_markdown or ["page"]).pop(0)
        else:
            number = int(re.search(r"SECTION: (\d+)\.", content).group(1)) if "SECTION:" in content else 0
            answer = answers.get(number, {})
            text = answer if isinstance(answer, str) else json.dumps(answer)
        return httpx.Response(200, json={"choices": [{"message": {"content": text}}]})

    return httpx.MockTransport(handler)


# What a VLM might plausibly return for a messy report (typos, 'Yes' notes, mixed marks)
RAW_ANSWERS = {
    1: {"name_first": "Shaik", "name_middle": "Salman", "father_name": "Ayub Khan", "age": "25 yrs",
        "category": {"selected": "Muslim", "evidence": "written"}},
    3: {
        "sex": {"selected": "Male: YES", "evidence": "tick"},
        "built": {"selected": "Stout", "evidence": "tick"},
        "height": {"selected": "5.6-6.0ft", "evidence": "bold"},
        "hair": {"selected": "Back combed", "evidence": "tick"},
        "mustaches": {"selected": "Turnee up", "evidence": "tick"},
        "face": {"selected": ["Round", "Oval"], "evidence": "tick"},
        "ear": "Normal",
        "identification_marks": "A Mole on the front side below neck",
    },
    4: {"occupation": {"selected": "Farmer", "evidence": "bold"},
        "marital_status": {"selected": "Unmarried", "evidence": "tick"}},
    5: {"family_stay_together": {"selected": "Yes", "evidence": "bold"},
        "members": [{"relation": "mother", "name": "Najiya Begum", "alive_dead": {"selected": "Alive"}},
                    {"relation": "Father", "name": "Ayub Khan", "criminal_background": {"selected": "No"}}]},
    9: {"habits": {"selected": ["Liquor", "Smoking", "Chewing Pan/Gutka/Ganja"], "evidence": "bold"}},
    15: {"cases": [{"ps": "Indalwai", "cr_no": "54/2026", "section_of_law": "8(C) r/w 27(b) of NDPS Act"},
                   {"ps": None, "cr_no": None}]},
    16: {"date_of_arrest": "06.04.2026", "crime_no_section": "54/2026", "police_station": "Indalwai"},
}


RAW_BY_KEY = {SECTION_BY_NUMBER[n].key: v for n, v in RAW_ANSWERS.items()}


# ---------------------------------------------------------------------------
# schema
# ---------------------------------------------------------------------------

def test_schema_has_32_numbered_sections_in_order():
    numbered = [s.number for s in SECTIONS if s.number]
    assert numbered == list(range(1, 33))
    keys = [s.key for s in SECTIONS]
    assert len(keys) == len(set(keys))


def test_schema_choice_fields_have_distinct_options():
    for s in SECTIONS:
        for f in list(s.fields) + [c for t in s.tables for c in t.columns]:
            if f.kind != "text":
                assert f.options, f"{s.key}.{f.key} has no options"
                assert len(set(o.lower() for o in f.options)) == len(f.options)


def test_physical_features_has_template_columns_1_to_20():
    spec = next(s for s in SECTIONS if s.number == 3)
    assert [f.column_no for f in spec.fields] == list(range(1, 21))


# ---------------------------------------------------------------------------
# validation: snapping, flags, shape
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    "raw,options,expected",
    [
        ("Fe-male", ["Male", "Female"], "Female"),
        ("Male: YES", ["Male", "Female"], "Male"),
        ("Turnee up", ["Small", "Long", "Turned Up"], "Turned Up"),
        ("Hunch / Back", ["Stout", "Hunch Back"], "Hunch Back"),
        ("Back combed", ["Long", "Back Combed"], "Back Combed"),
        ("6.6-7.0 ft: Yeas", ["6.1-6.5 ft", "6.6-7.0 ft"], "6.6-7.0 ft"),
        ("Muslim", ["SC", "ST", "BC", "Others"], None),
        ("long", ["Long hair style", "Short"], None),  # 'long' must not match a longer, different option
    ],
)
def test_snap_option(raw, options, expected):
    assert snap_option(raw, options) == expected


def test_blank_standard_has_every_field_and_no_values():
    std = blank_standard()
    assert list(std["sections"]) == [s.key for s in SECTIONS]
    phys = std["sections"]["s03_physical_features"]
    assert phys["sex"] == {"selected": None, "evidence": None, "other": None}
    assert phys["ear"] is None
    assert [r["relation"] for r in std["sections"]["s05_family_history"]["members"]] == list(FAMILY_RELATIONS)
    assert len(FAMILY_RELATIONS) == 15 and "Father's Sister" in FAMILY_RELATIONS  # the standard's 15 rows
    # a table with no data stays in the template as one blank row (not removed)
    shelter = std["sections"]["s12_shelter"]["shelter"]
    assert len(shelter) == 1 and all(v is None for v in shelter[0].values())


def test_build_standard_snaps_and_flags():
    std = build_standard(RAW_BY_KEY, not_found=["s28_source_of_procurement"])
    sec = std["sections"]
    assert sec["s03_physical_features"]["sex"] == {"selected": "Male", "evidence": "tick", "other": None}
    assert sec["s03_physical_features"]["height"]["selected"] == "5.6-6.0 ft"
    assert sec["s03_physical_features"]["hair"]["selected"] == "Back Combed"
    assert sec["s03_physical_features"]["mustaches"]["selected"] == "Turned Up"
    assert sec["s04_socio_economic_profile"]["occupation"]["selected"] == "Farmer"
    assert sec["s09_regular_habits"]["habits"]["selected"] == ["Liquor", "Smoking", "Chewing Pan/Gutka/Ganja"]
    # fixed-row table: always the standard's relations in order, model's row order and case ignored
    members = sec["s05_family_history"]["members"]
    assert [m["relation"] for m in members] == list(FAMILY_RELATIONS)
    assert members[0]["criminal_background"]["selected"] == "No"
    assert members[1]["name"] == "Najiya Begum" and members[1]["alive_dead"]["selected"] == "Alive"
    # an all-empty row of a free table is dropped
    assert len(sec["s15_cases_confessed"]["cases"]) == 1

    codes = {(f["code"], f["path"]) for f in std["review_flags"]}
    assert ("multiple_selected", "s03_physical_features.face") in codes
    assert std["sections"]["s03_physical_features"]["face"]["selected"] == "Round"
    assert ("unrecognized_option", "s01_individual_particulars.category") in codes
    assert sec["s01_individual_particulars"]["category"]["selected"] is None
    assert sec["s01_individual_particulars"]["category"]["other"] == "Muslim"
    assert ("section_not_found", "s28_source_of_procurement") in codes


def test_required_and_format_checks_flag_but_never_rewrite():
    std = build_standard({
        "s01_individual_particulars": {"name_first": "A", "age": "250"},
        "s02_unique_identifiers": {"aadhaar_number": "1234 5678", "mobile_number": "63006338"},
        "s16_arrest_particulars": {"date_of_arrest": "yesterday", "crime_no_section": "1/2026",
                                   "police_station": "X"},
    })
    codes = {(f["code"], f["path"]) for f in std["review_flags"]}
    assert ("missing_required", "s01_individual_particulars.father_name") in codes
    assert ("implausible_value", "s01_individual_particulars.age") in codes
    assert ("invalid_format", "s02_unique_identifiers.aadhaar_number") in codes
    assert ("invalid_format", "s02_unique_identifiers.mobile_number") in codes
    assert ("invalid_format", "s16_arrest_particulars.date_of_arrest") in codes
    assert std["sections"]["s01_individual_particulars"]["age"] == "250"  # value kept as read
    assert std["sections"]["s16_arrest_particulars"]["date_of_arrest"] == "yesterday"


def test_nil_values_are_kept_but_do_not_satisfy_required():
    std = build_standard({"s01_individual_particulars": {"father_name": "Nil"}})
    assert std["sections"]["s01_individual_particulars"]["father_name"] == "Nil"
    assert ("missing_required", "s01_individual_particulars.father_name") in {
        (f["code"], f["path"]) for f in std["review_flags"]}


# ---------------------------------------------------------------------------
# DOCX writer / reader / sectioner agree
# ---------------------------------------------------------------------------

def _write(tmp_path: Path, standard=None) -> Path:
    out = tmp_path / "ir.docx"
    render_docx(standard, out)
    return out


def test_blank_template_round_trips_to_32_sections(tmp_path):
    md = docx_to_markdown(_write(tmp_path))
    sections, missing = split_sections(md)
    assert missing == []
    assert len(sections) == 33  # 32 numbered + narrative
    # every option of the template is listed, none selected
    phys = sections["s03_physical_features"]
    assert "☐ Male" in phys and "☐ Female" in phys
    assert "☑" not in phys


def test_filled_report_shows_selected_options_with_tick_and_bold(tmp_path):
    std = build_standard(RAW_BY_KEY)
    md = docx_to_markdown(_write(tmp_path, std))
    phys = split_sections(md)[0]["s03_physical_features"]
    assert "**☑ Male**" in phys
    assert "**☑ 5.6-6.0 ft**" in phys
    assert "☐ Female" in phys  # unselected options keep an empty box and plain text
    assert "**☐" not in phys


def test_reader_keeps_wingdings_ticks_bold_and_bullets(tmp_path):
    """A report in the style of the uploaded samples: tick symbol, bold choice, bulleted options."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    t = doc.add_table(rows=1, cols=3)
    p = t.rows[0].cells[0].paragraphs[0]
    sym = OxmlElement("w:sym")
    sym.set(qn("w:font"), "Wingdings 2")
    sym.set(qn("w:char"), "F050")
    p.add_run()._r.append(sym)
    p.add_run("Male: YES")
    t.rows[0].cells[1].paragraphs[0].add_run("5.6-6.0 ft").bold = True
    t.rows[0].cells[2].paragraphs[0].add_run("Cat")
    path = tmp_path / "style.docx"
    doc.save(path)
    md = docx_to_markdown(path)
    assert "☑Male: YES" in md
    assert "**5.6-6.0 ft**" in md
    assert "| Cat |" in md or "Cat" in md


def test_sectioner_matches_titles_when_numbers_are_missing():
    md = "\n".join([
        "Facts of the case ...",
        "INDIVIDUAL PARTICULARS", "| Name | x |",
        "PHYSICAL FEATURES", "| Sex | Male |",
        "A long narrative line that merely mentions physical features in passing and goes on and on and on forever.",
        "FAMILY HISTORY", "| Father | y |",
    ])
    sections, missing = split_sections(md)
    assert "Name" in sections["s01_individual_particulars"]
    assert "Sex" in sections["s03_physical_features"] and "Father" not in sections["s03_physical_features"]
    assert "Father" in sections["s05_family_history"]
    assert "s28_source_of_procurement" in missing


def test_narrative_is_parsed_without_the_model():
    facts1 = "The facts of the case are that " + "x" * 220
    facts3 = "The facts of the case are that again " + "y" * 220
    n = parse_narrative("\n".join([
        facts1, "Accused: 1) A S/o B", "2) C S/o D", "Date:06.04.2026",
        "Sub-Inspector of Police,", "Indalwai,Nizamabad", facts3, "Basing on the above content I registered a case",
    ]))
    assert n["part1_facts"] == facts1 and n["part3_facts"] == facts3
    assert n["report_date"] == "06.04.2026"
    assert n["officer"] == "Sub-Inspector of Police" and n["officer_station"] == "Indalwai,Nizamabad"
    assert n["part1_accused"].startswith("Accused")
    assert n["part3_registration"].startswith("Basing on")


# ---------------------------------------------------------------------------
# VLM client + mapper (fake local server)
# ---------------------------------------------------------------------------

def test_extract_json_tolerates_fences_and_chatter():
    assert extract_json('Sure!\n```json\n{"a": {"b": "}"}}\n```\nDone') == {"a": {"b": "}"}}
    assert extract_json("no json here") is None
    assert extract_json('{"broken": ') is None


def test_client_refuses_non_local_endpoint(monkeypatch):
    _vlm_env(monkeypatch)
    monkeypatch.setenv("LLM_BASE_URL", "https://api.example.com/v1")
    get_settings.cache_clear()
    with pytest.raises(VlmError, match="local"):
        VlmClient(get_settings())


def test_client_requires_enabled_vlm(monkeypatch):
    monkeypatch.setenv("LLM_ENABLED", "false")
    get_settings.cache_clear()
    with pytest.raises(VlmError, match="disabled"):
        VlmClient(get_settings())


def test_mapper_maps_each_section_and_flags_missing_ones(monkeypatch, tmp_path):
    _vlm_env(monkeypatch)
    calls = []
    client = VlmClient(get_settings(), transport=fake_server(RAW_ANSWERS, calls=calls))
    md = docx_to_markdown(_write(tmp_path))
    std = map_markdown(md, client)
    # the blank template marks nothing, so the model's "Male" has no basis in the report: left empty, flagged
    assert std["sections"]["s03_physical_features"]["sex"]["selected"] is None
    assert any(f["code"] == "unsupported_selection" and f["path"].endswith("sex") for f in std["review_flags"])
    assert len(calls) == 32  # one request per section, none for the narrative
    # the model is told how selection is marked and which options exist
    prompt = next(c for c in calls if "SECTION: 3." in c)
    assert "☑" in prompt and "bold" in prompt and "Hunch Back" in prompt
    # blank template has no marks -> nothing hallucinated into sections the model left empty
    assert std["sections"]["s02_unique_identifiers"]["aadhaar_number"] is None


def test_mapper_retries_invalid_json_then_flags_failed_section(monkeypatch, tmp_path):
    _vlm_env(monkeypatch)
    answers = {1: "not json at all", 3: "still { not json", **{n: {} for n in range(2, 33) if n not in (1, 3)}}
    client = VlmClient(get_settings(), transport=fake_server(answers))
    md = docx_to_markdown(_write(tmp_path))
    std = map_markdown(md, client)
    failed = {f["section"] for f in std["review_flags"] if f["code"] == "mapping_failed"}
    assert failed == {"s01_individual_particulars", "s03_physical_features"}
    # a failed section is empty, not invented, and the others are unaffected
    assert std["sections"]["s03_physical_features"]["sex"]["selected"] is None


def test_server_down_raises_vlm_error(monkeypatch):
    _vlm_env(monkeypatch)

    def boom(request):
        raise httpx.ConnectError("refused")

    client = VlmClient(get_settings(), transport=httpx.MockTransport(boom))
    assert client.is_up() is False
    with pytest.raises(VlmError, match="unreachable"):
        client.complete("s", "u")


# ---------------------------------------------------------------------------
# flat view + corrections
# ---------------------------------------------------------------------------

def test_flatten_ids_and_corrections_snap_to_template_options():
    std = build_standard(RAW_BY_KEY)
    flat = {f.id: f for f in flatten(std)}
    assert flat["s03_physical_features.height"].value == "5.6-6.0 ft"
    assert flat["s03_physical_features.height"].serial == "(3)"
    assert flat["s05_family_history.members.0.name"].label.endswith("Father – Name")
    assert flat["s09_regular_habits.habits"].value == "Liquor, Smoking, Chewing Pan/Gutka/Ganja"
    assert flat["s03_physical_features.face"].flagged  # multiple_selected flag carries over

    fixed = apply_overrides(std, {
        "s03_physical_features.sex": "Fe-male",
        "s03_physical_features.ear": "Large",
        "s09_regular_habits.habits": "gambling; toddy",
        "s05_family_history.members.2.name": "Ravi",
    })
    sec = fixed["sections"]
    assert sec["s03_physical_features"]["sex"]["selected"] == "Female"
    assert sec["s03_physical_features"]["ear"] == "Large"
    assert sec["s09_regular_habits"]["habits"]["selected"] == ["Gambling", "Toddy"]
    assert sec["s05_family_history"]["members"][2]["name"] == "Ravi"
    assert std["sections"]["s03_physical_features"]["sex"]["selected"] == "Male"  # original untouched
    assert set_value(fixed, "s99_nope.x", "1") is False


# ---------------------------------------------------------------------------
# end to end through the API (upload DOCX -> VLM mapping -> JSON / DOCX / corrections)
# ---------------------------------------------------------------------------

@pytest.fixture()
def standard_client(make_client, monkeypatch):
    _vlm_env(monkeypatch)
    import app.services.extraction_service as es

    class Fake(es.IrStandardService):
        def client(self):
            if self._client is None:
                self._client = VlmClient(self.settings, transport=fake_server(RAW_ANSWERS))
            return self._client

    monkeypatch.setattr(es, "IrStandardService", Fake)
    with make_client(enable_auth=False) as c:
        yield c


def _upload(client, path: Path):
    r = client.post("/api/v1/documents/upload", files={"file": (path.name, path.read_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert r.status_code in (200, 201, 202), r.text
    return r.json()["id"] if "id" in r.json() else r.json()["document_id"]


def test_api_flow_standard_json_docx_and_corrections(standard_client, tmp_path):
    client = standard_client
    doc_id = _upload(client, _write(tmp_path, build_standard(RAW_BY_KEY)))

    res = client.get(f"/api/v1/documents/{doc_id}/results").json()["results"]
    assert res["extraction_status"] in ("complete", "partial")
    assert res["standard_ir"]["sections"]["s03_physical_features"]["height"]["selected"] == "5.6-6.0 ft"
    ids = {f["id"]: f for f in res["fields"]}
    assert ids["s03_physical_features.sex"]["value"] == "Male"
    assert ids["s03_physical_features.sex"]["page"] == 3  # UI groups by section number
    assert any(s["stage"] == "standard_template_mapping" for s in res["stage_results"])

    std = client.get(f"/api/v1/documents/{doc_id}/standard").json()
    assert list(std["sections"]) == [s.key for s in SECTIONS]

    # reviewer corrects a choice with a messy value; the standard JSON and DOCX follow
    r = client.patch(f"/api/v1/documents/{doc_id}/fields", json={"corrections": [
        {"field_id": "s03_physical_features.sex", "corrected_value": "Fe-male", "review_status": "corrected"}]})
    assert r.status_code == 200, r.text
    std = client.get(f"/api/v1/documents/{doc_id}/standard").json()
    assert std["sections"]["s03_physical_features"]["sex"]["selected"] == "Female"

    docx = client.get(f"/api/v1/documents/{doc_id}/standard.docx")
    assert docx.status_code == 200
    out = tmp_path / "download.docx"
    out.write_bytes(docx.content)
    phys = split_sections(docx_to_markdown(out))[0]["s03_physical_features"]
    assert "**☑ Female**" in phys and "☐ Male" in phys

    blank = client.get("/api/v1/standard-ir/template.docx")
    assert blank.status_code == 200 and blank.content[:2] == b"PK"


def test_scanned_pdf_pages_are_read_by_the_vlm_then_mapped(make_client, monkeypatch, tmp_path, fixtures_dir):
    """PDF path: each page image goes to the VLM, its markdown is sectioned and mapped; Paddle is not used."""
    _vlm_env(monkeypatch)
    import app.services.extraction_service as es

    page_md = docx_to_markdown(_write(tmp_path, build_standard(RAW_BY_KEY)))
    transcripts = [page_md] + ["(page intentionally blank)"] * 20
    calls = []

    class Fake(es.IrStandardService):
        def client(self):
            if self._client is None:
                self._client = VlmClient(self.settings, transport=fake_server(RAW_ANSWERS, transcripts, calls))
            return self._client

    monkeypatch.setattr(es, "IrStandardService", Fake)
    monkeypatch.setattr(es.OcrService, "recognize_png", lambda *a, **k: pytest.fail("Paddle must not run"))
    pdf = fixtures_dir / "sample_04_scanned_image.pdf"
    with make_client(enable_auth=False) as client:
        r = client.post("/api/v1/documents/upload", files={"file": (pdf.name, pdf.read_bytes(), "application/pdf")})
        doc_id = r.json().get("id") or r.json()["document_id"]
        res = client.get(f"/api/v1/documents/{doc_id}/results").json()["results"]
        assert res["standard_ir"]["sections"]["s03_physical_features"]["sex"]["selected"] == "Male"
        assert res["page_summaries"][0]["text_source"] == "vlm"
        assert "image" in calls  # the page image really was sent to the model
        raw = client.get(f"/api/v1/documents/{doc_id}/raw").json()
        assert raw["pages"][0]["lines"] and raw["pages"][0]["text_source"] == "ocr"


def test_falls_back_to_ocr_pipeline_when_vlm_is_down(make_client, monkeypatch, tmp_path):
    _vlm_env(monkeypatch)
    import app.services.extraction_service as es

    class Down(es.IrStandardService):
        def ready(self):
            return "VLM server at http://127.0.0.1:9/v1 is not reachable"

    monkeypatch.setattr(es, "IrStandardService", Down)
    with make_client(enable_auth=False) as client:
        doc_id = _upload(client, _write(tmp_path, build_standard(RAW_BY_KEY)))
        res = client.get(f"/api/v1/documents/{doc_id}/results").json()["results"]
        assert res["standard_ir"] is None
        assert any(w.startswith("STANDARD_IR_FALLBACK") for w in res["extraction_warnings"])
        assert client.get(f"/api/v1/documents/{doc_id}/standard").status_code == 404
